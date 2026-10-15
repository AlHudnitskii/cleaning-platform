import io
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from src.api import export, import_tasks, locations, sync, tasks
from src.infrastructure.database.connection import session_scope
from src.infrastructure.database.models import Task, TaskStatusHistory
from tests.integration.conftest import call, multipart, payload, utc

pytestmark = pytest.mark.integration


async def test_create_task_rejects_assignee_from_other_country(account):
    admin = await account("admin", None)
    cleaner_nl = await account("cleaner", "NL")

    response = await call(
        tasks.create_task,
        "POST",
        headers=admin.headers,
        json_body={"title": "Clean hall", "country": "DE", "assigned_to": str(cleaner_nl.user.id)},
    )

    assert response.status_code == 400


async def test_create_task_returns_assignee_and_location_names(account, location_factory):
    manager = await account("manager", "DE")
    cleaner = await account("cleaner", "DE")
    germany = await location_factory("Germany", "DE", "country")
    berlin = await location_factory("Berlin", "DE", "city", germany)

    response = await call(
        tasks.create_task,
        "POST",
        headers=manager.headers,
        json_body={
            "title": "Clean hall",
            "country": "NL",
            "location_id": str(berlin.id),
            "assigned_to": str(cleaner.user.id),
        },
    )

    assert response.status_code == 201
    body = payload(response)
    assert body["country"] == "DE"
    assert body["assigned_to_email"] == cleaner.user.email
    assert body["location_path"] == ["Germany", "Berlin"]


async def test_pagination_is_stable_for_equal_timestamps(account, task_factory):
    admin = await account("admin", None)
    created_at = utc(2026, 1, 1, 12)
    created = {str((await task_factory(created_at=created_at)).id) for _ in range(3)}

    seen: list[str] = []
    for page in ("1", "2", "3"):
        response = await call(tasks.get_tasks, headers=admin.headers, params={"page": page, "limit": "1"})
        seen.extend(t["id"] for t in payload(response)["data"])

    assert sorted(seen) == sorted(created)


async def test_sort_by_priority_puts_urgent_first(account, task_factory):
    admin = await account("admin", None)
    for priority in ("low", "urgent", "normal", "high"):
        await task_factory(priority=priority)

    response = await call(tasks.get_tasks, headers=admin.headers, params={"sort": "priority"})

    assert [t["priority"] for t in payload(response)["data"]] == ["urgent", "high", "normal", "low"]


async def test_location_filter_includes_subtree(account, task_factory, location_factory):
    admin = await account("admin", None)
    germany = await location_factory("Germany", "DE", "country")
    berlin = await location_factory("Berlin", "DE", "city", germany)
    netherlands = await location_factory("Netherlands", "NL", "country")
    inside = await task_factory(location_id=berlin.id)
    await task_factory(country="NL", location_id=netherlands.id)

    response = await call(tasks.get_tasks, headers=admin.headers, params={"location_id": str(germany.id)})

    assert [t["id"] for t in payload(response)["data"]] == [str(inside.id)]


async def test_quality_review_requires_completed_task(account, task_factory):
    manager = await account("manager", "DE")
    task = await task_factory(status="on_review")

    response = await call(
        tasks.review_quality, "POST", headers=manager.headers, route={"task_id": str(task.id)}, json_body={"score": 5}
    )

    assert response.status_code == 409


async def test_sync_applies_valid_changes_and_reports_each_failure(account, task_factory):
    cleaner = await account("cleaner", "DE")
    own = await task_factory(assigned_to=cleaner.user.id)
    second = await task_factory(assigned_to=cleaner.user.id)
    foreign = await task_factory()

    response = await call(
        sync.batch_sync,
        "POST",
        headers=cleaner.headers,
        json_body={
            "changes": [
                {"type": "status_change", "task_id": str(foreign.id), "status": "in_progress"},
                {"type": "status_change", "task_id": str(own.id), "status": "in_progress"},
                {"type": "status_change", "task_id": "not-a-uuid", "status": "in_progress"},
                {"type": "status_change", "task_id": str(second.id), "status": "completed"},
            ]
        },
    )

    body = payload(response)
    assert body["synced"] == 1
    assert [e["task_id"] for e in body["errors"]] == [str(foreign.id), "not-a-uuid", str(second.id)]
    async with session_scope() as session:
        assert (await session.get_one(Task, own.id)).status == "in_progress"
        assert (await session.get_one(Task, second.id)).status == "pending"


async def test_sync_rejects_change_older_than_server_state(account, task_factory):
    cleaner = await account("cleaner", "DE")
    task = await task_factory(assigned_to=cleaner.user.id, status="in_progress")
    server_change = datetime.now(UTC)
    async with session_scope() as session, session.begin():
        session.add(
            TaskStatusHistory(task_id=task.id, old_status="pending", new_status="in_progress", changed_at=server_change)
        )

    response = await call(
        sync.batch_sync,
        "POST",
        headers=cleaner.headers,
        json_body={
            "changes": [
                {
                    "type": "status_change",
                    "task_id": str(task.id),
                    "status": "on_hold",
                    "timestamp": (server_change - timedelta(minutes=5)).isoformat(),
                }
            ]
        },
    )

    body = payload(response)
    assert body["synced"] == 0
    assert "changed on the server" in body["errors"][0]["error"]


async def test_import_treats_empty_cells_as_missing(account):
    manager = await account("manager", "DE")
    csv = b"title,country,description,priority,rrule\nClean lobby,NL,,,\nVacuum,DE,Carpets,high,FREQ=DAILY\n"
    body, headers = multipart("file", "tasks.csv", csv, "text/csv")

    response = await call(import_tasks.import_tasks, "POST", headers={**headers, **manager.headers}, body=body)

    assert response.status_code == 200
    assert payload(response)["imported"] == 2
    async with session_scope() as session:
        rows = (await session.scalars(select(Task).order_by(Task.title))).all()
    first = next(t for t in rows if t.title == "Clean lobby")
    assert (first.country, first.description, first.rrule, first.is_recurring) == ("DE", None, None, False)
    second = next(t for t in rows if t.title == "Vacuum")
    assert (second.priority, second.rrule, second.is_recurring) == ("high", "FREQ=DAILY", True)


async def test_import_rejects_assignee_from_other_country(account):
    manager = await account("manager", "DE")
    cleaner_nl = await account("cleaner", "NL")
    csv = f"title,country,assignee_email\nClean lobby,DE,{cleaner_nl.user.email}\n".encode()
    body, headers = multipart("file", "tasks.csv", csv, "text/csv")

    response = await call(import_tasks.import_tasks, "POST", headers={**headers, **manager.headers}, body=body)

    assert response.status_code == 200
    assert payload(response)["imported"] == 0
    assert payload(response)["errors"] == ["Row 2: assignee must be an active cleaner from the task's country"]


async def test_excel_export_neutralizes_formulas(account, task_factory, monkeypatch):
    admin = await account("admin", None)
    await task_factory(title='=HYPERLINK("http://evil","click")', created_at=utc(2026, 3, 1))

    async def fake_upload(*args, **kwargs):
        return "https://blob/exports/x"

    monkeypatch.setattr(export, "upload_blob", fake_upload)

    response = await call(
        export.export_tasks,
        headers=admin.headers,
        params={"date_from": "2026-03-01", "date_to": "2026-03-01", "format": "excel"},
    )

    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.get_body()))["Tasks"]
    headers = [c.value for c in sheet[1]]
    assert sheet.cell(row=2, column=headers.index("title") + 1).value.startswith("'=")


async def test_location_hierarchy_is_validated(account, location_factory):
    admin = await account("admin", None)
    floor = await location_factory("Floor 1", "DE", "floor")

    response = await call(
        locations.create_location,
        "POST",
        headers=admin.headers,
        json_body={"name": "Berlin", "country": "DE", "level": "city", "parent_id": str(floor.id)},
    )

    assert response.status_code == 400


async def test_location_names_with_dots_do_not_break_paths(account, location_factory):
    admin = await account("admin", None)
    germany = await location_factory("Germany", "DE", "country")

    response = await call(
        locations.create_location,
        "POST",
        headers=admin.headers,
        json_body={"name": "St. Peter's Office №1", "country": "DE", "level": "building", "parent_id": str(germany.id)},
    )

    assert response.status_code == 201
    body = payload(response)
    assert body["full_name"] == "Germany / St. Peter's Office №1"
    assert body["path"].count(".") == 1


async def test_locations_are_scoped_to_manager_country(account, location_factory):
    manager = await account("manager", "DE")
    await location_factory("Germany", "DE", "country")
    await location_factory("Netherlands", "NL", "country")

    response = await call(locations.get_locations, headers=manager.headers, params={"country": "NL"})

    assert [loc["name"] for loc in payload(response)] == ["Germany"]
