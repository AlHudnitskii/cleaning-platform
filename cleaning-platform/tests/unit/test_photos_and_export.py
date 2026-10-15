import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.domain.services.photos import JPEG, PNG, WEBP, detect_image_type, display_filename
from src.infrastructure.export.exporter import export_to_excel, neutralize_formula, tasks_to_dataframe


@pytest.mark.parametrize(
    ("data", "want"),
    [
        (b"\xff\xd8\xff\xe0rest", JPEG),
        (b"\x89PNG\r\n\x1a\nrest", PNG),
        (b"RIFF\x00\x00\x00\x00WEBPVP8 ", WEBP),
        (b"<svg xmlns='http://www.w3.org/2000/svg'/>", None),
        (b"GIF89a", None),
        (b"", None),
    ],
    ids=["jpeg", "png", "webp", "svg rejected", "gif rejected", "empty rejected"],
)
def test_detect_image_type(data: bytes, want: object) -> None:
    assert detect_image_type(data) == want


@pytest.mark.parametrize(
    ("filename", "want"),
    [
        ("../../etc/passwd", "passwd"),
        ("C:\\Users\\me\\photo.jpg", "photo.jpg"),
        ("<script>.png", "_script_.png"),
        (None, "photo.jpg"),
        ("...", "photo.jpg"),
    ],
)
def test_display_filename_is_sanitized(filename: str | None, want: str) -> None:
    assert display_filename(filename, JPEG) == want


@pytest.mark.parametrize(
    ("value", "want"),
    [
        ('=HYPERLINK("http://evil","x")', '\'=HYPERLINK("http://evil","x")'),
        ("+1", "'+1"),
        ("-2", "'-2"),
        ("@SUM(A1)", "'@SUM(A1)"),
        ("Clean lobby", "Clean lobby"),
        (5, 5),
    ],
)
def test_neutralize_formula(value: object, want: object) -> None:
    assert neutralize_formula(value) == want


def test_excel_export_stores_formulas_as_text() -> None:
    df = tasks_to_dataframe([{"id": "1", "title": "=1+1", "description": "ok"}])
    workbook = load_workbook(io.BytesIO(export_to_excel(df)))
    sheet = workbook["Tasks"]
    headers = [cell.value for cell in sheet[1]]
    title = sheet.cell(row=2, column=headers.index("title") + 1)
    assert title.data_type == "s"
    assert title.value == "'=1+1"


def test_empty_export_keeps_columns() -> None:
    df = tasks_to_dataframe([])
    assert list(df.columns)[:2] == ["id", "title"]
    workbook = load_workbook(io.BytesIO(export_to_excel(df)))
    assert workbook["Tasks"].max_row == 1


def test_dataframe_has_fixed_column_order() -> None:
    df = tasks_to_dataframe([{"title": "b", "id": "a"}])
    assert isinstance(df, pd.DataFrame)
    assert list(df.columns)[:2] == ["id", "title"]
